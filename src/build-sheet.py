#!/usr/bin/env python3
"""Generates dist/comparison.xlsx from data/research.json.

Uploaded to Drive with conversion, it becomes the Google Sheet comparison table.
No formulas — it is a reference table, so there is nothing to recalculate.

Sheet 1 "Comparison"  the shortlist, in the order it was asked for
Sheet 2 "All vendors" every vendor in the dataset, same columns
Sheet 3 "Legend"      what each column means, and the sources behind each row
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = json.load(open(os.path.join(ROOT, "data", "research.json")))

# Rise tokens (kept in step with src/theme.js)
RISE, INK, MUTED, BORDER_C, FILL2 = "022BBE", "1B1B1F", "5B6070", "ECEDF1", "F7F8FA"
EDGE_FILL = {"technology": "E6ECFB", "relationships": "FDEFE5", "both": "EEF0F4"}
EDGE_TEXT = {"technology": "022BBE", "relationships": "B85C00", "both": "3A3F4B"}

# the shortlist, in the order the comparison was requested
SHORTLIST = ["cloudx", "mobilefuse", "liftoff-dsp", "inmobi", "digital-turbine",
             "amazon-aps", "pubmatic-openwrap", "magnite", "nimbus"]

COLUMNS = [
    ("Vendor",                      22),
    ("Category",                    19),
    ("Value proposition",           52),
    ("What the SDK is based on",    52),
    ("Market share / scale",        44),
    ("Niche (who it is for)",       38),
    ("Demand running on the traffic", 46),
    ("Edge: tech or relationships", 34),
    ("Pricing / take rate",         38),
    ("Watch-out",                   46),
]

HEAD_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Arial", size=10, color=INK)
BODY_FONT = Font(name="Arial", size=9.5, color=MUTED)
NAME_FONT = Font(name="Arial", size=10, bold=True, color=INK)
THIN = Side(style="thin", color=BORDER_C)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_WRAP = Alignment(wrap_text=True, vertical="top")


def sdk(i):
    for s in DATA["sdks"]:
        if s["id"] == i:
            return s
    raise SystemExit("unknown sdk id: " + i)


def based_on(s):
    """Integration model plus whatever is open about the stack."""
    bits = [s["integrationModel"]]
    o = s.get("openStack") or {}
    flags = []
    if o.get("prebid"):
        flags.append("Prebid Mobile path")
    if o.get("openSource"):
        flags.append("publishes source")
    if flags:
        bits.append("OPEN STACK: " + " + ".join(flags) + ". " + (o.get("note") or ""))
    else:
        bits.append("OPEN STACK: proprietary — no Prebid path and no published source.")
    if s.get("sdkSize") and "not publicly disclosed" not in s["sdkSize"].lower():
        bits.append("Footprint: " + s["sdkSize"])
    return "\n\n".join(b.strip() for b in bits if b and b.strip())


def edge_cell(s):
    e = s.get("edge") or {}
    return (e.get("basis", "").capitalize() or "—") + "\n" + (e.get("note") or "")


def write_table(ws, rows, title):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=13, bold=True, color=INK)
    ws["A2"] = ("Snapshot " + DATA["meta"]["snapshotDate"] +
                " · generated from data/research.json · " + DATA["meta"]["owner"])
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="8A8F98")
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 14

    head = 4
    for c, (label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=head, column=c, value=label)
        cell.font = HEAD_FONT
        cell.fill = PatternFill("solid", fgColor=RISE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[head].height = 30

    for r, s in enumerate(rows, start=head + 1):
        values = [
            s["name"], s["category"], s["coreDifferentiator"], based_on(s),
            s["marketShare"], s["positioning"], s["demandNetwork"],
            edge_cell(s), s["pricing"], s.get("structuralWeakness", ""),
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = TOP_WRAP
            cell.border = BOX
            cell.font = NAME_FONT if c == 1 else (CELL_FONT if c == 2 else BODY_FONT)
            if (r - head) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=FILL2)
        basis = (s.get("edge") or {}).get("basis")
        if basis in EDGE_FILL:
            ec = ws.cell(row=r, column=8)
            ec.fill = PatternFill("solid", fgColor=EDGE_FILL[basis])
            ec.font = Font(name="Arial", size=9.5, color=EDGE_TEXT[basis])
        ws.row_dimensions[r].height = 118

    ws.freeze_panes = ws.cell(row=head + 1, column=2)
    ws.auto_filter.ref = (f"A{head}:{get_column_letter(len(COLUMNS))}{head + len(rows)}")


wb = Workbook()

ws1 = wb.active
ws1.title = "Comparison"
write_table(ws1, [sdk(i) for i in SHORTLIST], "In-app SDK comparison — requested shortlist")

ws2 = wb.create_sheet("All vendors")
ordered = sorted(DATA["sdks"], key=lambda s: (s["category"], s["name"]))
write_table(ws2, ordered, "In-app SDK comparison — every vendor tracked (%d)" % len(ordered))

# ---- legend ----
ws3 = wb.create_sheet("Legend")
ws3.sheet_view.showGridLines = False
ws3["A1"] = "How to read this"
ws3["A1"].font = Font(name="Arial", size=13, bold=True, color=INK)
ws3.column_dimensions["A"].width = 30
ws3.column_dimensions["B"].width = 110

notes = [
    ("Value proposition", "The one thing the vendor does that its competitors do not. Taken from the vendor's own positioning, then cut back to what the sources support."),
    ("What the SDK is based on", "How demand actually reaches the app, plus whether any part of the stack is open. \"Prebid Mobile path\" means the vendor ships or documents a Prebid integration — not merely that it bids into someone else's Prebid auction."),
    ("Market share / scale", "Scale figures are rarely comparable across these companies: a mediation share, a device-reach number and a demand-side adapter install count measure three different things. Where no independent figure exists, the cell says so."),
    ("Niche", "The buyer or publisher the vendor is actually built for, which is usually narrower than its marketing claims."),
    ("Demand running on the traffic", "Whose budgets flow through the integration — own-and-operated demand, a partner list, or an open exchange."),
    ("Edge: tech or relationships", "Whether the offering wins on technology that a competitor would have to rebuild, or on commercial access that a competitor would have to negotiate. \"Both\" means the two reinforce each other and neither alone explains the position."),
    ("Pricing / take rate", "What the vendor itself publishes. Only %d of %d vendors publish a take rate at all; everything else in circulation is third-party estimate." % (
        sum(1 for s in DATA["sdks"] if s["disclosure"]["takeRate"]), len(DATA["sdks"]))),
    ("Watch-out", "The structural weakness — the thing that would undercut the pitch, which is usually absent from the vendor's own materials."),
    ("", ""),
    ("Caveat", "Vendor-stated figures are not audited. Conversion lifts, eCPM gains and adoption claims are as reported by the vendor. Several fields were flagged unresolved in the source pass and say so in the cell."),
    ("Regenerating", "This sheet is generated from data/research.json. Edit the data and run ./build.sh — do not edit the sheet by hand, it will be overwritten."),
]
r = 3
for k, v in notes:
    ws3.cell(row=r, column=1, value=k).font = Font(name="Arial", size=10, bold=True, color=RISE)
    c = ws3.cell(row=r, column=2, value=v)
    c.font = Font(name="Arial", size=10, color=MUTED)
    c.alignment = TOP_WRAP
    ws3.row_dimensions[r].height = 42 if v else 10
    r += 1

r += 1
ws3.cell(row=r, column=1, value="Sources by vendor").font = Font(name="Arial", size=11, bold=True, color=INK)
r += 1
for s in sorted(DATA["sdks"], key=lambda x: x["name"]):
    ws3.cell(row=r, column=1, value=s["name"]).font = Font(name="Arial", size=9.5, bold=True, color=INK)
    c = ws3.cell(row=r, column=2, value=s["sources"])
    c.font = Font(name="Arial", size=9, color=MUTED)
    c.alignment = TOP_WRAP
    ws3.row_dimensions[r].height = 26
    r += 1

out = os.path.join(ROOT, "dist", "comparison.xlsx")
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)
print("sheet           %d shortlist · %d all vendors  →  dist/comparison.xlsx"
      % (len(SHORTLIST), len(ordered)))
